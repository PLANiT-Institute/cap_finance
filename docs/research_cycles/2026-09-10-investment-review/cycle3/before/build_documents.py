from pathlib import Path
import re,json,csv,hashlib,math
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path('/Users/jinsu/Documents/GitHub/cap_finance')
OUT=ROOT/'output/pilot_review'; FIG=OUT/'figures'; FIG.mkdir(exist_ok=True)
DAT=OUT/'figure_data'; DAT.mkdir(exist_ok=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':'#202020','text.color':'#202020','axes.titleweight':'bold','savefig.dpi':220})
BLUE='#245475'; TEAL='#23847c'; RED='#b15b40'; GRAY='#727b83'

def rows(p):
    with open(ROOT/p,newline='') as f:return list(csv.DictReader(f))
ncc=rows('data/research_v2/ncc_timing_frontier/risk_metrics.csv')
qual=rows('data/research_v2/ncc_timing_frontier/qualification.csv')
steel=rows('data/research_v2/steel_procurement_tail/conditional_risk.csv')
curves=json.loads((ROOT/'data/research_v2/iea_2025_scenarios/curves.json').read_text())
assert len(ncc)==6 and len(qual)==16
assert sum(r['annual_feasible']=='True' for r in qual)==0
assert sum(r['eligible_for_cumulative_terminal_menu']=='True' for r in qual)==4
assert all(abs(float(r['theoretical_mean'])-3971.2314147902393)<1e-7 for r in steel)

def save(fig,name):
    fig.savefig(FIG/f'{name}.png',bbox_inches='tight',facecolor='white'); fig.savefig(FIG/f'{name}.pdf',bbox_inches='tight')
    from matplotlib.text import Text
    krfont=font_manager.FontProperties(fname='/System/Library/Fonts/Supplemental/AppleGothic.ttf')
    translations={'Crude oil':'원유','Japan LNG':'일본 LNG','Price ratio   2024 = 1':'가격비율   2024년 = 1','Discount weights':'할인 가중치','PV of one unit paid in year t':'해당 연도 비용 1의 현재가치','Assumed electricity emissions':'가상 전력 배출경로','tCO₂ per MWh':'tCO₂ / MWh','0.2 to 0':'0.2 → 0','0.4 to 0':'0.4 → 0','Grid 0.2 to 0':'전력계수 0.2 → 0','Grid 0.4 to 0':'전력계수 0.4 → 0','75% cumulative cap':'누적 상한 75%','Cumulative emissions / reference (%)':'기준안 대비 누적배출 (%)','Commissioning year':'가동 시작 연도','Pool 1 ES95':'표본 1 ES95','Pool 2 ES95':'표본 2 ES95','Common theoretical mean':'공통 이론 평균','Absolute cost  (million source EUR)':'절대비용 (백만 원문 EUR)','PLANiT energy system\nPrices and supply limits\nElectricity emissions':'플랜잇 에너지 시스템\n가격과 공급 제약\n전력 배출계수','Plant decision module\nTechnology and timing\nProcurement cash flows':'공장 의사결정 모듈\n기술과 투자 시점\n구매 현금흐름','Eligible plan comparison\nExpected cost and ES\nIndependent evaluation':'적격 계획 비교\n기대비용과 ES\n독립 평가','Public technology DB\nSource specifications\nVersioned transformations':'공개 감축기술 DB\n출처와 공정 사양\n판본별 변환 기록','Process quantities':'공정 물량','Later feedback\nDemand and investment\nSystem re-solving required':'후속 피드백\n수요와 투자 정보\n시스템 재계산 필요'}
    for tx in fig.findobj(Text):
        val=tx.get_text()
        if val in translations:tx.set_text(translations[val]);tx.set_fontproperties(krfont)
    fig.savefig(FIG/f'{name}_ko.png',bbox_inches='tight',facecolor='white');plt.close(fig)
def export(name,records):
    with open(DAT/name,'w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=records[0].keys());w.writeheader();w.writerows(records)

fig,axs=plt.subplots(1,2,figsize=(9,2.7),layout='constrained')
scenario_records=[]
for ax,commodity,title in zip(axs,['crude_oil','gas_japan'],['Crude oil','Japan LNG']):
    found=[c for c in curves if c['commodity']==commodity]
    if not found and commodity=='gas_japan':
        found=[c for c in curves if 'japan' in str(c['region']).lower() and 'gas' in str(c['commodity']).lower()]
    assert len(found)==3,(commodity,[c['commodity'] for c in curves])
    for c in found:
        color={'CPS':GRAY,'STEPS':BLUE,'NZE':TEAL}[c['scenario']]
        ratios=np.array(c['values'])/c['values'][0]
        ax.plot(c['years'],ratios,'o-',label=c['scenario'],color=color)
        for y,v,ratio in zip(c['years'],c['values'],ratios):scenario_records.append({'commodity':title,'scenario':c['scenario'],'year':y,'source_price':v,'ratio_2024':ratio})
    ax.set(title=title,xlim=(2024,2050),ylabel='Price ratio   2024 = 1'); ax.set_xticks([2024,2035,2050]);ax.grid(alpha=.15);ax.legend(frameon=False,fontsize=8,ncol=3,loc='lower left')
save(fig,'scenario');export('scenario_nodes.csv',scenario_records)

fig,axs=plt.subplots(1,2,figsize=(9,2.7),layout='constrained'); years=np.arange(2027,2051)
dis=[]
for r,color in zip([.03,.05,.07],[GRAY,BLUE,TEAL]):
    fac=(1+r)**(-(years-2027));axs[0].plot(years,fac,color=color,label=f'{r:.0%}')
    for y,v in zip(years,fac):dis.append({'year':int(y),'discount_rate':r,'discount_factor':v})
axs[0].set(title='Discount weights',ylabel='PV of one unit paid in year t');axs[0].legend(frameon=False,ncol=3)
for g,color in zip([.2,.4],[TEAL,BLUE]):axs[1].plot(years,g*(2050-years)/23,label=f'{g} to 0',color=color)
axs[1].set(title='Assumed electricity emissions',ylabel='tCO₂ per MWh');axs[1].legend(frameon=False)
for ax in axs:ax.set_xticks([2027,2035,2040,2050]);ax.grid(alpha=.15)
save(fig,'discount_grid');export('discount_weights.csv',dis)

fig,ax=plt.subplots(figsize=(9,2.65),layout='constrained'); x=np.arange(4)
for i,(cond,color,label) in enumerate([('g0.2_r0.0',TEAL,'Grid 0.2 to 0'),('g0.4_r0.0',BLUE,'Grid 0.4 to 0')]):
    rr=[r for r in qual if r['condition']==cond]; vals=[float(r['cumulative_ratio'])*100 for r in rr]
    b=ax.bar(x+(i-.5)*.34,vals,.32,color=color,label=label)
    ax.bar_label(b,labels=[f'{v:.2f}%' for v in vals],padding=3,fontsize=9)
ax.axhline(75,ls='--',color=RED,lw=1.3,label='75% cumulative cap');ax.set_xticks(x,['2030','2035','2040','2045']);ax.set(ylabel='Cumulative emissions / reference (%)',xlabel='Commissioning year',ylim=(0,105));ax.legend(frameon=False,ncol=3,loc='upper center',fontsize=9);ax.grid(axis='y',alpha=.12)
save(fig,'ncc_eligibility');export('ncc_qualification.csv',qual);export('ncc_risk_metrics.csv',ncc)

fig,ax=plt.subplots(figsize=(9,2.9),layout='constrained')
for i,(seed,color) in enumerate([('20260923',BLUE),('20260924',TEAL)]):
    rr=[r for r in steel if r['seed']==seed and r['pairing']=='joint'];rr.sort(key=lambda r:float(r['beta']))
    b=ax.bar(np.arange(3)+(i-.5)*.3,[float(r['es_alpha']) for r in rr],.28,color=color,label=f'Pool {i+1} ES95')
    ax.bar_label(b,labels=[f"{float(r['es_alpha']):,.0f}" for r in rr],padding=3,fontsize=9)
ax.axhline(3971.2314,ls='--',color=GRAY,label='Common theoretical mean');ax.set_xticks([0,1,2],['β = 0','β = 0.5','β = 1']);ax.set(ylabel='Absolute cost  (million source EUR)',ylim=(0,11800));ax.legend(frameon=False,ncol=3,fontsize=9,loc='upper center');ax.grid(axis='y',alpha=.12)
save(fig,'steel');export('steel_risk_metrics.csv',[r for r in steel if r['pairing']=='joint'])

fig,ax=plt.subplots(figsize=(9,3.15));ax.set(xlim=(0,10),ylim=(0,4));ax.axis('off')
boxes=[(.1,2.3,2.7,1.35,'PLANiT energy system\nPrices and supply limits\nElectricity emissions'),(3.65,2.3,2.7,1.35,'Plant decision module\nTechnology and timing\nProcurement cash flows'),(7.2,2.3,2.7,1.35,'Eligible plan comparison\nExpected cost and ES\nIndependent evaluation'),(3.65,.05,2.7,1.05,'Public technology DB\nSource specifications\nVersioned transformations')]
for x,y,w,h,t in boxes:
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04',edgecolor='#657c8a',facecolor='#f2f6f8',lw=1));ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=10,linespacing=1.4)
for x1,x2 in [(2.83,3.62),(6.38,7.17)]:ax.annotate('',xy=(x2,3),xytext=(x1,3),arrowprops=dict(arrowstyle='->',color=BLUE,lw=1.7))
ax.annotate('',xy=(5,2.26),xytext=(5,1.14),arrowprops=dict(arrowstyle='->',color=BLUE,lw=1.7));ax.text(5.13,1.64,'Process quantities',fontsize=9)
ax.plot([4.2,4.2,1.45],[2.25,1.8,1.8],color=TEAL,lw=1.5,ls='--'); ax.annotate('',xy=(1.45,2.25),xytext=(1.45,1.8),arrowprops=dict(arrowstyle='->',color=TEAL,lw=1.5,ls='--'))
ax.text(.4,1.05,'Later feedback\nDemand and investment\nSystem re-solving required',fontsize=9,color=TEAL)
save(fig,'integration')


# A quantile display exposes TCaR without inventing a probability density.
for tcar_lang in ['en','ko']:
    fig,ax=plt.subplots(figsize=(9,2.8),layout='constrained')
    for row_i,year in enumerate([2030,2035,2040]):
        rr=[r for r in ncc if r['candidate']==f'replace_{year}']
        for k,r in enumerate(rr):
            y=2-row_i + (.07 if k==0 else -.07)
            lo,hi,es=[float(r[v]) for v in ['p50','q_alpha','es_alpha']]
            ax.plot([lo,hi],[y,y],color=BLUE if k==0 else TEAL,lw=2,ls='-' if k==0 else '--')
            ax.scatter([lo],[y],c=BLUE if k==0 else TEAL,marker='o',s=28)
            ax.scatter([hi],[y],c=BLUE if k==0 else TEAL,marker='s',s=28)
            ax.scatter([es],[y],c=BLUE if k==0 else TEAL,marker='D',s=26)
        ax.text((float(rr[0]['p50'])+float(rr[0]['q_alpha']))/2,2-row_i+.21,'TCaR95  '+ ' / '.join(f"{float(r['tcar']):.1f}" for r in rr),ha='center',fontsize=9)
    ax.set(yticks=[2,1,0],yticklabels=['2030','2035','2040'],ylim=(-.4,2.6),xlim=(1700,5900))
    ax.set_xlabel('Incremental cost  (million EUR, assumed 2019 operations)' if tcar_lang=='en' else '증분비용  (백만 EUR, 운영가격 2019년 가정)')
    ax.set_ylabel('Commissioning' if tcar_lang=='en' else '가동 시작')
    ax.set_title('Circle P50   Square P95   Diamond ES95   |   Solid pool 1   Dashed pool 2' if tcar_lang=='en' else '원 P50   사각 P95   마름모 ES95   |   실선 표본 1   점선 표본 2',fontsize=10)
    ax.grid(axis='x',alpha=.15)
    if tcar_lang=='ko':
        kf=font_manager.FontProperties(fname='/System/Library/Fonts/Supplemental/AppleGothic.ttf')
        for tx in [ax.xaxis.label,ax.yaxis.label,ax.title]:tx.set_fontproperties(kf)
    fig.savefig(FIG/f"tcar{'_ko' if tcar_lang=='ko' else ''}.png",bbox_inches='tight',facecolor='white');plt.close(fig)
export('ncc_tcar_quantiles.csv',ncc)

TABLES={
'assumptions':{
'en':(['Input','Frozen company case','Process reference experiments'],[
['Horizon and discount','2027–2050; 5%','2027–2050; 5%; three construction years'],
['Scenario source','NZ15 and B20 local proxies','IEA 2025 ratios; separate assumed grid paths'],
['Prices and samples','GBM type; 10,000 paths; α = 0.90','Joint monthly blocks; 2 × 5,000 paths; α = 0.95'],
['Volatility inputs','Annual log SD: electricity 0.24245; hydrogen 0.25; CAPEX 0.06','Demeaned WB 2010–2025 log changes; 12 month blocks'],
['Physical output','Fixed; 22 allocated assets','NCC 0.92 Mt ethylene/year; steel 0.75 Mt liquid steel/year'],
['Emissions requirement','Cumulative Scope 1 cap; annual failures reported','NCC cumulative ≤75% of reference and terminal zero'],
['Key assumptions','Energy residual shares; replacement; no support','Grid 0.2/0.4 to zero; residual 0/0.05; pellet β 0/0.5/1'],
['Monetary interpretation','Declared 2025 KRW; source bases not fully verified','NCC assumes EUR2019 operations; steel retains source EUR'],
]),
'ko':(['입력','동결 기업 사례','후속 공정 참조'],[
['기간·할인율','2027–2050년; 5%','2027–2050년; 5%; 건설 3년'],
['시나리오 출처','로컬 NZ15·B20 대용치','IEA 2025 비율; 별도 가상 전력경로'],
['가격·표본','GBM형; 10,000경로; 90%','공동 월별 블록; 2 × 5,000경로; 95%'],
['변동성 입력','연간 로그 표준편차: 전력 0.24245; 수소 0.25; CAPEX 0.06','WB 2010–2025 로그 변화량; 12개월 블록'],
['생산 서비스','고정생산; 추정 자산 22개','NCC 에틸렌 92만 톤/년; 철강 액강 75만 톤/년'],
['배출 조건','누적 Scope 1 상한; 연간 위반 별도','NCC 기준안 누적의 75% 이하 및 말기 0'],
['핵심 가정','구입 에너지 잔존비율; 교체투자; 지원 0','전력 0.2/0.4→0; 잔여 0/0.05; β 0/0.5/1'],
['화폐값 해석','2025 KRW 선언; 원천 기준 일부 미검증','NCC 운영가격 EUR2019 가정; 철강 원문 EUR'],
])},
'choices':{'en':(['Company and scenario','Mean minimum','TCaR90 minimum','ES90 minimum'],[]),'ko':(['기업·시나리오','평균 최소','TCaR90 최소','ES90 최소'],[])},
'ncc':{'en':(['Start','Mean','TCaR95 pool 1','TCaR95 pool 2','ES95 pool 1','ES95 pool 2'],[]),'ko':(['가동','평균','TCaR95 표본 1','TCaR95 표본 2','ES95 표본 1','ES95 표본 2'],[])},
'validation':{
'en':(['Validation question','Current evidence','Remaining requirement'],[
['Source and units','Selected primary observations and transformations registered','Complete plant scope and harmonised monetary bases'],
['Physical and cash arithmetic','Separate quantity and cash checks; explicit exclusions','Measured site balances and procurement settlement'],
['Conditional sampling','Independent pools and paired bootstrap','Selection uncertainty and wider candidate exploration'],
['Model specification','Price generator and process sensitivity comparisons','Long horizon distribution identification'],
['External decision validity','Conditional changes in contracts and eligibility','Matched corporate case and independent partner review']]),
'ko':(['검증 질문','현재 증거','남은 확인'],[
['출처·단위','선택 원문값과 변환 등록','완결된 공장 범위·화폐 기준 통일'],
['물리·현금 산술','물량·현금 별도 대조와 제외항목 명시','현장 수지와 실제 구매·결제'],
['조건부 표본','독립 표본·공통 인덱스 부트스트랩','선택오차와 넓은 후보 탐색'],
['모형 사양','가격생성법·공정 민감도 비교','장기 미래분포 식별'],
['외부 의사결정','조건부 계약 선택·자격 변화','실제 기업 매핑·독립 협력팀 검토']])},
'db':{
'en':(['Record group','Required content'],[
['Identity and applicability','Technology version, region, vintage, output product, scale and system boundary'],
['Physical coefficients','Feedstock grade, yield, external purchases, internal energy, coproduct use and sales'],
['Emissions qualification','Direct and electricity factors, residual sources, capture versus avoided emissions'],
['Capital and operation','Price year, currency, CAPEX scope, construction time, component lives, replacement and downtime'],
['Evidence and transformation','Original value, unit, source location, conversion formula, status, alternatives and exclusions'],
['Release and reproducibility','Redistribution status, version history, machine readable schema and executable reference case']]),
'ko':(['레코드 구분','필수 내용'],[
['정체성·적용 범위','기술 판본·지역·연도·제품·규모·시스템 경계'],
['물리 계수','원료 등급·수율·외부 구매·내부 에너지·부산물 사용과 판매'],
['배출 자격','직접·전력 배출계수·잔여 배출·포집량과 회피량 구분'],
['투자·가동','가격연도·통화·CAPEX 범위·건설기간·구성 수명·교체·중단'],
['증거·변환','원문값·단위·출처 위치·변환식·자료 상태·대안·제외 이유'],
['공개·재현','재배포 상태·변경 이력·기계 판독 구조·실행 참조 사례']])},
'next':{
'en':(['Experiment','Evidence needed to proceed'],[
['Matched process comparison','Same service and boundaries; reconciled purchases, internal flows and emissions'],
['PLANiT scenario exchange','Jointly consistent prices, supply and electricity emissions; recorded model version'],
['Procurement and settlement','A defensible mapping from benchmarks to invoices; priced feasible contracts'],
['Decision and uncertainty','Independent mean–ES comparison; menu expansion; separated sampling and specification effects'],
['Public replication','External team reproduces a nonconfidential case and documents corrections']]),
'ko':(['실험','다음 단계로 갈 증거'],[
['같은 경계의 공정 비교','같은 서비스·경계; 구매·내부 흐름·배출의 수지 대조'],
['플랜잇 시나리오 교환','가격·공급·전력 배출의 정합성 및 모형 판본 기록'],
['구매·결제 반영','벤치마크와 청구가격의 연결 근거; 실제 비용을 가진 계약'],
['선택·불확실성 평가','독립 평균–ES 비교; 후보 확장; 표본·사양 영향 분리'],
['공개 재현','외부 팀의 비기밀 사례 재현과 오류 수정 기록']])}}
choices=[['LOTTE · B20','0% / 0','100% / 0','100% / 0'],['LOTTE · NZ15','0% / 0','75% / 1','75% / 0'],['Mitsui · B20','0% / 0','100% / 0','75% / 0'],['Mitsui · NZ15','0% / 0','100% / 0','100% / 0'],['Nippon Steel · B20','0% / 0','100% / 0','75% / 0'],['Nippon Steel · NZ15','0% / 0','100% / 0','100% / 0'],['POSCO · B20','0% / 0','75% / 1','75% / 0'],['POSCO · NZ15','0% / 0','100% / 1','75% / 0']]
for lang in ['en','ko']:
    TABLES['choices'][lang][1].extend(choices)
    for year in [2030,2035,2040]:
        rr=[r for r in ncc if r['candidate']==f'replace_{year}'];TABLES['ncc'][lang][1].append([str(year),f"{float(rr[0]['analytic_mean']):,.2f}",f"{float(rr[0]['tcar']):,.2f}",f"{float(rr[1]['tcar']):,.2f}",f"{float(rr[0]['es_alpha']):,.2f}",f"{float(rr[1]['es_alpha']):,.2f}"])

# Native editable Office Math with linear notation, kept deliberately compact.
EQS={'loss':'L(x, ω | s, θ) = C(x, ω | s, θ) − C(x₀, ω | s, θ)     (1)',
'cash':'C(x, ω) = ∑ₜ dₜ [Kₜ + Oₜ + ∑ⱼ Pⱼₜqⱼₜ + Zₜ + Hₜ − Sₜ] − dₜᵣVₜᵣ     (2)',
'risk':'μ(x) = E[L(x)]       TCaRα(x) = qα[L(x)] − q₀.₅[L(x)]     (3)',
'es':'ESα(L) = minη {η + E[max(L − η, 0)] / (1 − α)}     (4)'}
# T denotes terminal date; use a simple final subscript rather than ambiguous tr.
EQS['cash']='C(x, ω) = ∑ₜ dₜ [Kₜ + Oₜ + ∑ⱼ Pⱼₜqⱼₜ + Zₜ + Hₜ − Sₜ] − d(T)V(T)     (2)'

def run_font(run,font,size=None):
    run.font.name=font
    if size:run.font.size=Pt(size)
    pr=run._element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),font)

def inline(p,txt,font):
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*)',txt):
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m:
            rel=p.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245475');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        elif part:
            r=p.add_run(part.strip('*') if part.startswith('**') else part);run_font(r,font);r.bold=True if part.startswith('**') else (None if font=='NanumGothic' else False)

def table(doc,key,lang,font):
    headers,data=TABLES[key][lang];t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={'assumptions':[1.15,2.5,3.05],'choices':[2.05,1.5,1.65,1.5],'ncc':[.6,1.22,1.22,1.22,1.22,1.22],'validation':[1.6,2.45,2.65],'db':[1.7,5.0],'next':[1.8,4.9]}[key]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate([headers]+data):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trPr=cells[0]._tc.getparent().get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
        if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            tcPr.append(mar);bd=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            tcPr.append(bd);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'245475' if i==0 else ('F3F6F8' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            if j>0 and key in ['choices','ncc']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(txt);run_font(r,font,9.3 if lang=='en' else 9.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '111111')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def mathrun(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r

def sub(base,index):
    e=OxmlElement('m:sSub');b=OxmlElement('m:e');b.append(mathrun(base));ix=OxmlElement('m:sub');ix.append(mathrun(index));e.append(b);e.append(ix);return e

def formula(key):
    m=OxmlElement('m:oMath')
    if key=='loss':parts=['L(x, ω | s, θ) = C(x, ω | s, θ) − C(',sub('x','0'),', ω | s, θ)   (1)']
    elif key=='cash':parts=['C(x, ω) = ',sub('∑','t'),sub('d','t'),' [',sub('K','t'),' + ',sub('O','t'),' + ',sub('∑','j'),sub('P','jt'),sub('q','jt'),' + ',sub('Z','t'),' + ',sub('H','t'),' − ',sub('S','t'),'] − ',sub('d','T'),sub('V','T'),'   (2)']
    elif key=='risk':parts=['μ(x) = E[L(x)];     ',sub('TCaR','α'),'(x) = ',sub('q','α'),'[L(x)] − ',sub('q','0.5'),'[L(x)]   (3)']
    else:parts=[sub('ES','α'),'(L) = ',sub('min','η'),' {η + E[max(L − η, 0)] / (1 − α)}   (4)']
    for part in parts:m.append(mathrun(part) if isinstance(part,str) else part)
    return m

for lang,source,filename in [('en','PAPER_EN.md','PLANiT_Transition_Pilot_EN.docx'),('ko','REVIEW_KO.md','PLANiT_Transition_Pilot_KO.docx')]:
    text=(ROOT/'paper/pilot_review'/source).read_text().replace('박진수 |','Jinsu Park |')
    (ROOT/'paper/pilot_review'/source).write_text(text)
    doc=Document()
    for bd in doc.styles.element.xpath('//w:pBdr'):
        bd.getparent().remove(bd)
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.9);sec.footer_distance=Inches(.3)
    font='Times New Roman' if lang=='en' else 'NanumGothic'
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[name];st.font.name=font;st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12 if lang=='en' else Pt(14.5);normal.paragraph_format.space_after=Pt(8)
    for name,size in [('Title',23),('Subtitle',13),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle';st.font.italic=False;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(8)
    doc.styles['Caption'].font.size=Pt(9.2);doc.styles['Caption'].paragraph_format.line_spacing=1.06
    doc.core_properties.author='Jinsu Park';doc.core_properties.title=text.splitlines()[0][2:];doc.core_properties.subject='Industrial transition pilot methodology';doc.core_properties.keywords='PLANiT, transition costs, Expected Shortfall, TCaR'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    for block in text.split('\n\n'):
        block=block.strip()
        if not block:continue
        if block=='\\pagebreak':doc.add_page_break();continue
        if block.startswith('{{table:'):table(doc,block[8:-2],lang,font);continue
        if block.startswith('{{fig:'):
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(FIG/f"{block[6:-2]}{'_ko' if lang=='ko' else ''}.png"),width=Inches(6.7));continue
        if block.startswith('{{eq:'):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(9)
            p._p.append(formula(block[5:-2]));continue
        if block.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,block[2:].replace(' 아래 산업',' 아래\n산업') if lang=='ko' else block[2:],font)
        elif block.startswith('## '):p=doc.add_paragraph(style='Subtitle');inline(p,block[3:],font)
        elif block.startswith('### '):p=doc.add_paragraph(style='Heading 1' if re.match(r'### (\d+ |References|참고|Appendix|부록)',block) else 'Heading 2');inline(p,block[4:],font)
        else:
            cap=bool(re.match(r'^(Figure [1-6]\.|Table [2-6]\.|그림 [1-6]\.|표 [2-6]\.)',block))
            p=doc.add_paragraph(style='Caption' if cap else 'Normal');inline(p,block,font)
            if block.startswith('['):p.paragraph_format.space_after=Pt(7)
    doc.save(OUT/filename)

inputs=['paper/methodology/TCaR_METHOD_KO.md','data/research_v2/iea_2025_scenarios/curves.json','data/research_v2/ncc_timing_frontier/risk_metrics.csv','data/research_v2/ncc_timing_frontier/qualification.csv','data/research_v2/steel_procurement_tail/conditional_risk.csv']
manifest={'basis':'TCaR v2.32 and separate saved reference experiments','new_corporate_simulation':False,'tables':6,'figures':6,'inputs':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in inputs},'outputs':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.glob('*.docx')}}
(OUT/'provenance.json').write_text(json.dumps(manifest,indent=2));print(json.dumps(manifest,indent=2))
