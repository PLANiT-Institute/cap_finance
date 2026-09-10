"""Read-only primary-workbook audit. Use bundled Python; never resave the source.
This checks selected physical arithmetic, not the complete Excel calculation engine.
"""
from pathlib import Path
import csv, hashlib, json, math, warnings
from datetime import datetime, timezone
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT/'tmp/pdfs/tcar_research_v2/langhorst2026'
OUT = ROOT/'data/research_v2/langhorst_plant_reference'
MODEL = RAW/'model.xlsx'
ARTICLE = 'https://doi.org/10.1007/s44498-026-00166-1'
DATA = 'https://doi.org/10.5281/zenodo.19094433'

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def write_csv(name, rows):
    with (OUT/name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    metadata = json.loads((RAW/'zenodo.json').read_text())
    expected = next(f['checksum'].split(':')[1] for f in metadata['files'] if f['key']=='Steel plant model.xlsx')
    assert hashlib.md5(MODEL.read_bytes()).hexdigest()==expected
    original = sha(MODEL)
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='DrawingML support is incomplete')
        wb = openpyxl.load_workbook(MODEL, data_only=True)
        formulas = openpyxl.load_workbook(MODEL, data_only=False)
    native=[]
    def val(sheet, cell):
        value=wb[sheet][cell].value
        formula=formulas[sheet][cell].value if formulas[sheet][cell].data_type=='f' else None
        if formula is not None and not isinstance(formula,str):
            formula={'text':formula.text,'ref':formula.ref,'type':'array'}
        native.append(dict(sheet=sheet,cell=cell,value=value,formula=formula))
        return value
    G='7.1 - Data - General'; R='1. Results'; checks=[]
    def check(name, result, expected, scope):
        ok=math.isclose(result,expected,rel_tol=1e-10,abs_tol=1e-8)
        checks.append(dict(check=name,reconstructed=result,source_cached=expected,absolute_error=abs(result-expected),passed=ok,scope=scope))
        assert ok,checks[-1]
    params=[]
    for i in [19,77,78,83,88,96,98,99,105,106,114,115,126,127,128,129,131,132,133,135,136,137,140,143,145,146,148]:
        params.append(dict(parameter=val(G,f'B{i}'),value=val(G,f'D{i}'),unit=val(G,f'E{i}'),source_reported_reference=val(G,f'F{i}'),sheet=G,cell=f'D{i}',status='source_model_input_not_independently_validated_upstream'))
    for cell in ['B3','B9']:
        params.append(dict(parameter='AEC LHV efficiency '+('2050' if cell=='B3' else '2025'),value=val('7.2 - Data - Electrolysis',cell),unit='MWh_H2_LHV/MWh_electricity',source_reported_reference='Danish Energy Agency: SI year 2024; workbook header 2025; upstream edition not resolved',sheet='7.2 - Data - Electrolysis',cell=cell,status='source_model_input_not_independently_validated_upstream'))
    g=lambda i:val(G,f'D{i}')
    plate=g(133); slab_ref=g(132); export=slab_ref-plate/g(114)
    rows=[]
    names=['Reference','H2-DRI-EAF','H2 reheating','EE.1','EE.2','ME.1','ME.2','All EE&ME']
    fe=[s for s in wb.sheetnames if s.startswith('5.') and 'total' not in s]
    rm=[s for s in wb.sheetnames if s.startswith('4.')]
    aec=[s for s in wb.sheetnames if s.startswith('6.')]
    for j,name in enumerate(names):
        energy_row=5+j; roll_row=17+j; emission_row=57+j
        # Aggregation of carriers; source '-' means absent in the particular source scenario.
        def carrier(cell):
            x=val(R,cell)
            assert isinstance(x,(int,float)) or x in ['-',None]
            return float(x) if isinstance(x,(int,float)) else 0.
        power=sum(carrier(f'{c}{energy_row}') for c in ['F','G','H','J'])+sum(carrier(f'{c}{roll_row}') for c in ['F','G'])
        check(name+' electricity aggregate',power,val(R,f'P{roll_row}'),'arithmetic over source component caches')
        gas=carrier(f'I{energy_row}')+carrier(f'E{roll_row}')
        coal=carrier(f'D{energy_row}'); pci=carrier(f'E{energy_row}'); cog=carrier(f'D{roll_row}')
        gross=val(R,f'M{roll_row}')
        check(name+' gross versus internal COG',coal+pci+power+gas+cog,gross,'source subsystem inflows, internal transfer counted')
        # Rebuild production quantities from final product vector and yields.
        improved=j in [5,7]
        rolling_in=plate/g(115 if improved else 114)
        slab=rolling_in+export; ls=slab/g(106 if improved else 105)
        sf=fe[j]; sr=rm[j]
        check(name+' plate output',plate,val(sr,'J78'),'fixed product')
        check(name+' exported slabs',export,val(sf,'H81' if j==0 else 'AD81'),'fixed coproduct')
        check(name+' liquid steel',ls,val(sf,'K66' if j==0 else 'Q52'),'yield reconstruction')
        check(name+' cast slabs',slab,val(R,f'AB{emission_row}'),'product denominator')
        if j==0:
            pellet=val(sf,'I31'); ext_scrap=val(sf,'G47'); internal_scrap=val(sf,'G46'); h2=0.; dri=0.
        else:
            ext_scrap=val(sf,'O24'); internal_scrap=val(sf,'O21')
            ef=g(96 if j in [6,7] else 98); sec=g(99)
            fe_out=ls*g(145)
            total_fe=fe_out/(ef*sec)
            dri=(total_fe-(ext_scrap+internal_scrap)*g(146))/g(137)
            pellet=dri*g(136)
            check(name+' DRI',dri,val(sf,'Q19'),'iron-yield and scrap reconstruction')
            check(name+' pellets',pellet,val(sf,'Q3'),'DRI input recipe')
            h2_dr=dri*g(135)/1000
            check(name+' DR hydrogen',h2_dr,val(sf,'N9'),'58 kg per tonne DRI source recipe')
            eta=val('7.2 - Data - Electrolysis','B3' if j in [3,7] else 'B9')
            el_dr=h2_dr*g(148)/eta
            check(name+' electrolyser electricity',el_dr,val(R,f'G{energy_row}'),'hydrogen LHV divided by efficiency')
            check(name+' EAF electricity',ls*g(83)/1000,val(sf,'O39'),'source EAF recipe')
            if j>=2:
                h2=val(R,f'G{83+j}')
                check(name+' total hydrogen',sum(val(R,f'{c}{83+j}') for c in ['D','E','F']),h2,'DR, reheating and EAF hydrogen')
            else:h2=h2_dr
        direct=val(R,f'D{emission_row}')+val(R,f'F{emission_row}')
        indirect=power*val(sr,'C8')
        check(name+' indirect emissions',indirect,val(R,f'E{emission_row}')+val(R,f'G{emission_row}'),'Swedish source coefficient, not KR/JP')
        check(name+' emissions aggregate',direct+indirect,val(R,f'H{emission_row}'),'source rounded process emissions retained')
        upstream=val(R,f'D{emission_row}')+val(R,f'E{emission_row}')
        check(name+' labelled slab intensity',upstream/slab,val(R,f'W{emission_row}'),'slab denominator, not liquid steel')
        rows.append(dict(scenario=name,plates_kt_y=plate,export_slabs_kt_y=export,liquid_steel_kt_y=ls,cast_slabs_kt_y=slab,pellets_kt_y=pellet,external_scrap_kt_y=ext_scrap,internal_scrap_kt_y=internal_scrap,dri_kt_y=dri,onsite_hydrogen_kt_y=h2,electricity_demand_before_CHP_GWh_y=power,natural_gas_input_GWh_y=gas,steelmaking_natural_gas_GWh_y=carrier(f'I{energy_row}'),rolling_natural_gas_GWh_y=carrier(f'E{roll_row}'),coking_coal_input_GWh_y=coal,PCI_coal_input_GWh_y=pci,internal_COG_to_rolling_GWh_y=cog,source_gross_subsystem_energy_GWh_y=gross,carrier_inputs_excluding_internal_COG_GWh_y=coal+pci+power+gas,potential_process_gas_export_GWh_y=-carrier(f'K{energy_row}'),potential_recovered_heat_GWh_y=-carrier(f'M{energy_row}')-carrier(f'H{roll_row}'),direct_emissions_kt_y=direct,source_indirect_emissions_kt_y=indirect,source_steelmaking_emissions_per_t_slab=upstream/slab,derived_steelmaking_emissions_per_t_liquid_steel=upstream/ls))
    check('COG replacement in rolling',rows[1]['natural_gas_input_GWh_y']-carrier('I6')-(rows[0]['natural_gas_input_GWh_y']-carrier('I5')),rows[0]['internal_COG_to_rolling_GWh_y'],'same heat service')
    native=list({(r['sheet'],r['cell']):r for r in native}.values())
    (OUT/'selected_native_cells.json').write_text(json.dumps(native,ensure_ascii=False,indent=2)+'\n')
    write_csv('scenario_physical_ledger.csv',rows);write_csv('parameter_register.csv',params);write_csv('arithmetic_checks.csv',checks)
    gaps=[
      dict(id='internal_energy',status='separate_before_FIN',detail='Source subsystem gross energy includes internal COG; never price it again as an external purchase.'),
      dict(id='CHP_boundary',status='hold_net_purchase',detail='Power demand is before CHP conversion. Exportable gases and recoverable heat are physical potential; no generation efficiency or sale contract is identified.'),
      dict(id='product_denominator',status='hold_crude_steel_claim',detail='Results W57:W64 divide iron/steel emissions by cast slabs AB57:AB64; preserve this instead of treating the article crude-steel wording as liquid-steel intensity.'),
      dict(id='pellet_grade',status='hold_generic_ore_mapping',detail='BF Fe grade input is 0.67; DRI recipe uses 1.42 t pellets/tDRI and DRI Fe 0.92, implying another pellet Fe fraction. Ore grade and pellet premiums need separate prices.'),
      dict(id='costs',status='hold_frontier',detail='No CAPEX, timing, supplier prices, utilization distribution or domestic financing identified in this physical study. Efficiency scenarios are not costless selectable investments.'),
      dict(id='source_edition',status='hold_upstream_validation',detail='Electrolyser source year differs: SI 2024 versus workbook 2025. Recorded values agree with article; original DEA edition needs verification.'),
      dict(id='thermodynamic_closure',status='not_claimed',detail='Energy layer tracks carrier inputs and selected recovered heat, not a closed first-law heat balance. Primary authors acknowledge baseline carbon imbalance.'),
      dict(id='corporate_transfer',status='hold_corporate_calibration',detail='Swedish public output plus literature model, not measured complete SSAB flows, POSCO, Nippon Steel or a GCAM forecast. No probabilities on eight deterministic cases.')]
    (OUT/'adoption_gates.json').write_text(json.dumps(gaps,ensure_ascii=False,indent=2)+'\n')
    assert sha(MODEL)==original
    report=dict(audited_at_utc=datetime.now(timezone.utc).isoformat(),article_doi=ARTICLE,dataset_doi=DATA,dataset_publication_date=metadata['metadata']['publication_date'],article_publication_date='2026-09-05',dataset_md5_matched=True,source_workbook_unchanged=True,source_sheets=len(wb.sheetnames),selected_source_cells=len(native),parameter_rows=len(params),scenarios=len(rows),arithmetic_checks=len(checks),checks_passed=all(c['passed'] for c in checks),native_excel_recalculated=False,full_model_recalculated=False,new_MC_runs=0,FIN_adoption=False,corporate_calibration=False,primary_PDF_pages_visually_reviewed=[6,7,10,11],sources={n:sha(RAW/n) for n in ['article.pdf','si.docx','model.xlsx','Readme.txt','zenodo.json']},outputs={p.name:sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='manifest.json'},builder_sha256=sha(Path(__file__)))
    (OUT/'manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['sources','outputs']},ensure_ascii=False))

if __name__=='__main__':main()
