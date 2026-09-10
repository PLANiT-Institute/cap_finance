"""Read DEA source caches without saving the workbook; no stochastic calibration.

Run with the bundled Python containing openpyxl. The optional plot is generated
separately from the resulting comparison CSV. No production inputs are changed.
"""
from pathlib import Path
import csv
import hashlib
import json
import math
from datetime import datetime, timezone
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'tmp/pdfs/tcar_research_v2/steel_upstream'
OUT = ROOT / 'data/research_v2/dea_electrolyser_reference'
SHEET = '80 AEC 100 MW'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def save_csv(name, rows):
    with (OUT / name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    original = sha(RAW / 'dea_aug2026.xlsx')
    wb = openpyxl.load_workbook(RAW / 'dea_aug2026.xlsx', data_only=True)
    formulas = openpyxl.load_workbook(RAW / 'dea_aug2026.xlsx', data_only=False)
    ws = wb[SHEET]
    assert wb['Intro']['D4'].value == '0015'
    assert wb['Intro']['D7'].value == 2025
    rows = [7,8,11,12,13,15,16,17,18,19,20,21,22,23,24,25,27,28,29,30,31,32,33,34,35,38,39,45,46,49,50,52]
    native=[]; parameters=[]; checks=[]
    for r in rows:
        for c in 'CDEFGHIJK':
            cell=f'{c}{r}'
            raw_formula=formulas[SHEET][cell]
            formula=raw_formula.value if raw_formula.data_type == 'f' else None
            native.append(dict(sheet=SHEET,cell=cell,cached_value=ws[cell].value,formula=formula))
            parameters.append(dict(parameter=ws[f'B{r}'].value,year=ws[f'{c}3'].value,
                source_estimate_label=ws[f'{c}4'].value,value=ws[cell].value,
                source_cell=cell,note_ids=ws[f'L{r}'].value,reference_ids=ws[f'M{r}'].value,
                status='DEA_2026_08_reference_not_company_calibration'))
    for r in [61,62,63,64,65,66,67,68,69,71,72,73,74,75,76,77]:
        native.append(dict(sheet=SHEET,cell=f'B{r}',cached_value=ws[f'B{r}'].value,formula=None))
    for cell in ['D4','D6','D7']:
        native.append(dict(sheet='Intro',cell=cell,cached_value=wb['Intro'][cell].value,formula=None))
    def ck(name, calculated, cached):
        ok=math.isclose(calculated,cached,rel_tol=1e-10,abs_tol=1e-8)
        checks.append(dict(check=name,calculated=calculated,source_cached=cached,passed=ok))
        assert ok, checks[-1]
    for c in 'CDEFGHIJK':
        v=lambda r: ws[f'{c}{r}'].value
        ck(c+' LHV',33.333/v(12)*100,v(15))
        ck(c+' HHV',39.388/v(12)*100,v(16))
        ck(c+' hydrogen yield',1000/v(12),v(20))
        ck(c+' full-load daily capacity',v(7)*1000*24/v(12),v(8))
        ck(c+' total heat',100-v(16),v(17))
        ck(c+' heat accounting',v(18)+v(19),v(17))
        # Component central estimates form an accounting sum. Component marginal
        # bounds need not occur together; do not add them as joint quantiles.
        if c in 'CDEFG':
            ck(c+' central CAPEX components',sum(v(r) for r in range(28,35)),v(27))
        ck(c+' CAPEX daily-output normalization',v(27)*v(7)*1000/v(8),v(35))
    comparisons=[]
    for year,col,eta in [(2025,'C',.587),(2050,'G',.6994)]:
        old=33.33/eta
        new=ws[f'{col}12'].value
        comparisons.append(dict(year=year,Langhorst_source_kWh_per_kg_H2=old,
            DEA_AEC100_reference_kWh_per_kg_H2=new,
            substitution_change_percent=(new/old-1)*100,
            hydrogen_recipe_kg_per_t_DRI=58,
            Langhorst_electrolysis_MWh_per_t_DRI=old*58/1000,
            DEA_electrolysis_MWh_per_t_DRI=new*58/1000,
            electricity_change_MWh_per_t_DRI=(new-old)*58/1000,
            status='fixed_58kg_recipe_algebra_only_not_matched_edition_or_plant_rerun'))
    interval_diagnostics=[]
    for year,c,l,u in [(2025,'C','H','I'),(2050,'G','J','K')]:
        low=ws[f'{l}27'].value; high=ws[f'{u}27'].value
        # Equal-tail 5/95 lognormal interpretation is a diagnostic assumption.
        # DEA supplies no joint distribution or identification of the central
        # estimate with the median; do not use this fit for simulation.
        median=math.sqrt(low*high)
        interval_diagnostics.append(dict(year=year,source_central=ws[f'{c}27'].value,
            source_lower=low,source_upper=high,
            assumed_equal_tail_lognormal_median=median,
            median_vs_source_central_percent=(median/ws[f'{c}27'].value-1)*100,
            adopted_for_simulation=False))
    save_csv('parameter_register.csv',parameters)
    save_csv('arithmetic_checks.csv',checks)
    save_csv('specification_comparison.csv',comparisons)
    save_csv('interval_identification.csv',interval_diagnostics)
    (OUT/'selected_native_cells.json').write_text(json.dumps(native,ensure_ascii=False,indent=2)+'\n')
    assert sha(RAW/'dea_aug2026.xlsx') == original
    manifest=dict(audited_at_utc=datetime.now(timezone.utc).isoformat(),edition='August 2026 / 0015',
        source_urls={'landing':'https://ens.dk/en/analyses-and-statistics/technology-data-renewable-fuels',
                     'catalogue':'https://ens.dk/media/7823/download',
                     'workbook':'https://ens.dk/media/6444/download',
                     'brief':'https://ens.dk/media/3352/download',
                     'midrex':'https://www.midrex.com/direct-from-midrex/technical-articles/hot-briquetted-iron-steels-most-versatile-metallic-in-the-transition-to-the-hydrogen-economy/'},
        source_sha256={p.name:sha(p) for p in [RAW/'dea_aug2026.pdf',RAW/'dea_aug2026.xlsx',RAW/'dea_brief_aug2026.pdf',RAW/'dea_page.html',RAW/'midrex2021.html']},
        selected_native_cells=len(native),parameter_observations=len(parameters),arithmetic_checks=len(checks),all_passed=True,
        source_workbook_unchanged=True,full_Excel_recalculation=False,Monte_Carlo_runs=0,
        source_stated_interval_coverage=.90,distribution_identified=False,company_calibrated=False,
        old_DEA_edition_resolved=False,Millner_full_paper_obtained=False,
        output_sha256={p.name:sha(p) for p in OUT.iterdir() if p.suffix in ['.csv','.json','.png'] and p.name!='manifest.json'})
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in manifest.items() if k not in ['source_urls','source_sha256','output_sha256']},indent=2))

if __name__=='__main__': main()
